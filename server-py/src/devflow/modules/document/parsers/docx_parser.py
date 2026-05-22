from __future__ import annotations

import io

from devflow.modules.document.parsers.base import ParsedBlock, ParsedDocument


class DocxParser:
    name = "docx"
    suffixes = (".docx",)

    def parse(self, blob: bytes, *, filename: str | None = None) -> ParsedDocument:
        from docx import Document as DocxDocument

        doc = DocxDocument(io.BytesIO(blob))
        blocks: list[ParsedBlock] = []
        title: str | None = None
        for p in doc.paragraphs:
            text = (p.text or "").strip()
            if not text:
                continue
            style = (p.style.name or "").lower() if p.style else ""
            if "heading" in style:
                level = 1
                for ch in style:
                    if ch.isdigit():
                        level = int(ch)
                        break
                if title is None and level == 1:
                    title = text
                blocks.append(ParsedBlock(text=text, kind="heading", level=level))
            else:
                blocks.append(ParsedBlock(text=text))
        return ParsedDocument(title=title or filename, blocks=blocks)
